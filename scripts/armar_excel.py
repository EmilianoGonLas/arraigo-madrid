# -*- coding: utf-8 -*-
"""
Arma el Excel final "Arraigo 2011" con la forma que pide la actividad:
filas = año de llegada a la vivienda, columnas = VUPACI.

    python3 armar_excel.py
"""

import os

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
ENTRADA = os.path.join(RAIZ, "salidas", "arraigo_por_vupaci.csv")
SECCIONADO = os.path.join(RAIZ, "datos", "originales", "Seccionado VUPACI.xlsx")
INDICADORES = os.path.join(RAIZ, "datos", "ine", "csv", "C2011_ccaa13_Indicadores.csv")
SALIDA = os.path.join(RAIZ, "salidas", "Arraigo 2011 - completo.xlsx")

# el mismo orden de filas que la Hoja1 del Excel original
FILAS = ["2011", "2010", "2009", "2008", "2007", "2006",
         "2001-2005", "1991-2000", "1981-1990", "1971-1980", "Antes de 1971"]

# la tabla del INE abre con el Total, antes de las categorías
FILAS_CON_TOTAL = ["Total"] + FILAS


def poblacion_por_vupaci():
    """
    Control independiente: población total de cada VUPACI según el fichero
    masivo de indicadores por sección censal del INE (tabla t1_1).

    No tiene por qué coincidir con el total de la consulta —aquello son
    personas residentes en viviendas principales y esto es población total,
    y cada cifra va redondeada por su lado— pero sirve para detectar un
    seccionado mal leído o un grupo que se haya consultado de más o de menos.
    """
    if not os.path.exists(INDICADORES):
        return None
    sec = pd.read_excel(SECCIONADO, header=1).iloc[:, 1:3]
    sec.columns = ["VUPACI", "SECCION"]
    sec = sec.dropna(how="any")
    sec["VUPACI"] = sec["VUPACI"].astype(int)
    sec["SECCION"] = sec["SECCION"].astype("int64").astype(str).str.zfill(10)

    ind = pd.read_csv(INDICADORES, dtype=str)
    ind["SECCION"] = (ind["cpro"].str.zfill(2) + ind["cmun"].str.zfill(3)
                      + ind["dist"].str.zfill(2) + ind["secc"].str.zfill(3))
    ind["t1_1"] = pd.to_numeric(ind["t1_1"], errors="coerce")

    unido = sec.merge(ind[["SECCION", "t1_1"]], on="SECCION", how="left")
    return unido.groupby("VUPACI").agg(
        poblacion_indicadores=("t1_1", "sum"),
        secciones_sin_datos=("t1_1", lambda x: int(x.isna().sum())),
    )


def _leer_marcas(df):
    """
    {vupaci: {categoría: motivo}} a partir de las columnas de trazabilidad.

    - `marcadas`: la cifra venía con asterisco del INE (por debajo del umbral
      de fiabilidad). Se concentran en las categorías viejas, donde la muestra
      del censo da menos apoyo: 46 de 66 son "Antes de 1971".
    - `ausentes`: el INE no dibujó esa columna y el 0 lo pusimos nosotros.
    """
    marcas = {}
    for v, txt in df["marcadas"].dropna().items():
        for cat in str(txt).split("|"):
            if cat:
                marcas.setdefault(v, {})[cat] = "asterisco"
    for v, txt in df["ausentes"].dropna().items():
        for cat in str(txt).split("|"):
            if cat:
                marcas.setdefault(v, {})[cat] = "ausente"
    return marcas


def _escribir_hoja1(libro, tabla, marcas):
    """
    Escribe la Hoja1 con la misma disposición que la plantilla original:
    fila 1 -> "VUPACI" y los códigos; fila 2 -> "Año de llegada a la vivienda";
    fila 3 -> el Total de cada VUPACI, como en la tabla del INE;
    de la fila 4 en adelante, un año de llegada por fila.

    Las cifras que el INE marcó con asterisco quedan en ámbar y con nota; los
    ceros que pusimos nosotros, en gris y en cursiva. El valor se guarda como
    número, no como texto, para que las columnas se puedan seguir sumando.
    """
    hoja = libro.create_sheet("Hoja1")
    negrita = Font(bold=True)
    relleno_total = PatternFill("solid", fgColor="E5E7EB")
    fuente_ast = Font(color="B45309", bold=True)
    relleno_ast = PatternFill("solid", fgColor="FEF3C7")
    fuente_cero = Font(color="6B7280", italic=True)

    hoja.cell(row=1, column=1, value="VUPACI").font = negrita
    for j, v in enumerate(tabla.columns, start=2):
        c = hoja.cell(row=1, column=j, value=int(v))
        c.font = negrita
        c.alignment = Alignment(horizontal="center")

    hoja.cell(row=2, column=1, value="Año de llegada a la vivienda").font = negrita

    n_ast = n_cero = 0
    for i, anyo in enumerate(tabla.index, start=3):
        etiqueta = hoja.cell(row=i, column=1, value=anyo)
        if anyo == "Total":
            etiqueta.font = negrita
            etiqueta.fill = relleno_total
        for j, v in enumerate(tabla.columns, start=2):
            valor = tabla.loc[anyo, v]
            celda = hoja.cell(row=i, column=j,
                              value=None if pd.isna(valor) else int(valor))
            if anyo == "Total":
                celda.font = negrita
                celda.fill = relleno_total
            motivo = marcas.get(v, {}).get(anyo)
            if motivo == "asterisco":
                celda.font = fuente_ast
                celda.fill = relleno_ast
                celda.comment = Comment(
                    "El INE publica esta cifra marcada con asterisco: queda por "
                    "debajo del umbral de fiabilidad de la muestra censal. "
                    "Úsese con cautela.", "Control")
                n_ast += 1
            elif motivo == "ausente":
                celda.font = fuente_cero
                celda.comment = Comment(
                    "El INE no publicó esta categoría para este VUPACI: la "
                    "columna no aparecía en la tabla. Es un cero (el total "
                    "cuadra con la suma de las demás categorías).", "Control")
                n_cero += 1

    fila = len(tabla.index) + 4
    hoja.cell(row=fila, column=1,
              value="Leyenda").font = negrita
    c = hoja.cell(row=fila + 1, column=1,
                  value=f"Ámbar ({n_ast} celdas): cifra marcada con asterisco "
                        f"por el INE — por debajo del umbral de fiabilidad de "
                        f"la muestra. Se concentran en las categorías más "
                        f"antiguas.")
    c.font = fuente_ast
    c = hoja.cell(row=fila + 2, column=1,
                  value=f"Gris cursiva ({n_cero} celdas): el INE no publicó la "
                        f"categoría para ese VUPACI; es un cero.")
    c.font = fuente_cero

    hoja.column_dimensions["A"].width = 28
    for j in range(2, len(tabla.columns) + 2):
        hoja.column_dimensions[get_column_letter(j)].width = 8
    hoja.freeze_panes = "B3"
    print(f"Hoja1: {n_ast} celdas con asterisco, {n_cero} ceros no publicados")
    return hoja


def main():
    df = pd.read_csv(ENTRADA)
    df = df.sort_values("VUPACI").set_index("VUPACI")

    # Hoja1: la tabla tal cual la pide la actividad (traspuesta),
    # con la fila de Total delante como en la tabla del INE
    tabla = df[FILAS_CON_TOTAL].T
    tabla.index.name = "Año de llegada a la vivienda"
    tabla.columns.name = "VUPACI"

    # comprobación: cada columna tiene que sumar el Total que dio el INE
    suma = df[FILAS].sum(axis=1)
    control = pd.DataFrame({
        "total_INE": df["Total"],
        "suma_categorias": suma,
        "diferencia": df["Total"] - suma,
        "n_secciones": df["n_secciones"],
    })
    descuadres = control[control["diferencia"] != 0]

    pob = poblacion_por_vupaci()
    if pob is not None:
        control = control.join(pob)
        control["dif_vs_indicadores"] = (control["total_INE"]
                                         - control["poblacion_indicadores"])

    with pd.ExcelWriter(SALIDA, engine="openpyxl") as xls:
        df.reset_index().to_excel(xls, sheet_name="Datos INE", index=False)
        control.reset_index().to_excel(xls, sheet_name="Control", index=False)
        _escribir_hoja1(xls.book, tabla, _leer_marcas(df))
        xls.book.move_sheet("Hoja1", offset=-2)

    print(f"VUPACI: {len(df)} | filas: {len(tabla)}")
    print(f"personas totales: {df['Total'].sum():,}".replace(",", "."))
    print(f"descuadres (suma != total del INE): {len(descuadres)}")
    if len(descuadres):
        print(descuadres.head(10).to_string())
    print(f"\nExcel: {SALIDA}")


if __name__ == "__main__":
    main()
