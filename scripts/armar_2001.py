# -*- coding: utf-8 -*-
"""
Arma el arraigo de 2001 por VUPACI, para los grupos en que se puede.

El seccionado VUPACI es de 2011 y las secciones censales se redibujaron entre
2001 y 2011 (Alcalá de Henares, por ejemplo, tenía 8 distritos en 2001 y tiene
5 en 2011), así que un código de sección de 2011 no designa el mismo territorio
en 2001. Repetir la extracción con esa lista daría números de otros barrios.

Sí se puede, en cambio, para los VUPACI que son **municipios enteros**: los
municipios de Madrid son los mismos 179 en los dos censos. Para esos grupos el
dato de 2001 es directamente comparable.

Entrada : salidas/censo2001_municipios_madrid.csv (una consulta al INE)
Salida  : salidas/Arraigo 2001 - municipios enteros.xlsx
"""

import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from descargar_arraigo import leer_seccionado

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
M2001 = os.path.join(RAIZ, "salidas", "censo2001_municipios_madrid.csv")
IND2011 = os.path.join(RAIZ, "datos", "ine", "csv", "C2011_ccaa13_Indicadores.csv")
ARRAIGO2011 = os.path.join(RAIZ, "salidas", "arraigo_por_vupaci.csv")
SALIDA = os.path.join(RAIZ, "salidas", "Arraigo 2001 - municipios enteros.xlsx")

CAT2001 = ["2001", "2000", "1999", "1998", "1997", "1996", "1991-1995",
           "1981-1990", "1971-1980", "1961-1970", "1941-1960", "Antes de 1941"]


def vupaci_de_municipios_enteros():
    """VUPACI cuyas secciones cubren por completo los municipios que tocan."""
    grupos = leer_seccionado()
    ind = pd.read_csv(IND2011, dtype=str)
    ind["S"] = (ind.cpro.str.zfill(2) + ind.cmun.str.zfill(3)
                + ind.dist.str.zfill(2) + ind.secc.str.zfill(3))
    por_mun = ind.groupby(ind.S.str[:5])["S"].apply(set).to_dict()

    enteros, parciales = {}, []
    for v, secs in grupos.items():
        muns = sorted({x[:5] for x in secs})
        if all(por_mun.get(m, set()) <= set(secs) for m in muns):
            enteros[v] = muns
        else:
            parciales.append(v)
    return enteros, parciales, grupos


def main():
    m01 = pd.read_csv(M2001, dtype={"cmun": str}).set_index("cmun")
    enteros, parciales, grupos = vupaci_de_municipios_enteros()

    filas = []
    for v, muns in sorted(enteros.items()):
        sub = m01.reindex(muns)
        if sub[CAT2001].isna().any().any():
            raise RuntimeError(f"faltan municipios de la VUPACI {v}: {muns}")
        fila = {"VUPACI": v, "n_municipios": len(muns),
                "municipios": " ".join(muns),
                "Total": int(sub["TOTAL"].sum())}
        fila.update({c: int(sub[c].sum()) for c in CAT2001})
        filas.append(fila)
    df = pd.DataFrame(filas).set_index("VUPACI")

    a11 = pd.read_csv(ARRAIGO2011).set_index("VUPACI")
    comparacion = pd.DataFrame({
        "total_2001": df["Total"],
        "total_2011": a11["Total"].reindex(df.index),
    })
    comparacion["variación"] = comparacion["total_2011"] - comparacion["total_2001"]

    with pd.ExcelWriter(SALIDA, engine="openpyxl") as xls:
        df.reset_index().to_excel(xls, sheet_name="Arraigo 2001", index=False)
        comparacion.reset_index().to_excel(xls, sheet_name="2001 vs 2011", index=False)
        pd.DataFrame({"VUPACI_sin_dato_2001": sorted(parciales)}).to_excel(
            xls, sheet_name="No comparables", index=False)
        _formatear(xls.book)

    print(f"VUPACI con dato de 2001: {len(df)} de {len(grupos)}")
    print(f"  personas 2001: {df['Total'].sum():,}".replace(",", "."))
    print(f"VUPACI sin dato (parten municipios): {len(parciales)}")
    print("Excel:", SALIDA)


def _formatear(libro):
    negrita = Font(bold=True)
    aviso = PatternFill("solid", fgColor="FEF3C7")
    for nombre in libro.sheetnames:
        hoja = libro[nombre]
        for c in hoja[1]:
            c.font = negrita
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for j in range(1, hoja.max_column + 1):
            hoja.column_dimensions[get_column_letter(j)].width = 13
        hoja.freeze_panes = "B2"
    hoja = libro["No comparables"]
    hoja.column_dimensions["A"].width = 24
    hoja.cell(row=1, column=3,
              value="Estos VUPACI parten municipios, así que hacen falta "
                    "secciones. Las secciones de 2001 no son las de 2011 "
                    "(el seccionado se redibujó), por lo que el dato de 2001 "
                    "no se puede obtener con la lista de códigos actual.").font = negrita
    hoja.cell(row=1, column=3).fill = aviso
    hoja.column_dimensions["C"].width = 100


if __name__ == "__main__":
    main()
