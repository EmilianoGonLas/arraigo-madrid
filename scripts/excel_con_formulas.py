# -*- coding: utf-8 -*-
"""
Libro de Excel con el índice de arraigo calculado **con fórmulas vivas**:
cada celda derivada es una fórmula que se puede seguir y auditar, y los cuatro
pesos están en celdas sueltas, así que cambiar 0,33 por otra cosa recalcula
las 489 filas solo.

Hojas:
  Cálculo  las once categorías del INE + los cuatro tramos + proporciones + índice
  Pesos    los cuatro pesos, editables
  Método   qué se decidió y por qué

    python3 excel_con_formulas.py
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
ENTRADA = os.path.join(RAIZ, "salidas", "arraigo_por_vupaci.csv")
SALIDA = os.path.join(RAIZ, "salidas", "Indice de arraigo 2011 - con formulas.xlsx")

CAT = ["2011", "2010", "2009", "2008", "2007", "2006",
       "2001-2005", "1991-2000", "1981-1990", "1971-1980", "Antes de 1971"]

NEGRITA = Font(bold=True)
AMBAR_F = Font(color="B45309", bold=True)
AMBAR_R = PatternFill("solid", fgColor="FEF3C7")
GRIS = Font(color="6B7280", italic=True)
CABECERA = PatternFill("solid", fgColor="E5E7EB")


def main():
    df = pd.read_csv(ENTRADA).set_index("VUPACI").sort_index()
    libro = Workbook()

    pesos = libro.active
    pesos.title = "Pesos"
    _hoja_pesos(pesos)

    hoja = libro.create_sheet("Cálculo")
    _hoja_calculo(hoja, df)

    _hoja_metodo(libro.create_sheet("Método"))
    libro.move_sheet("Cálculo", offset=-1)
    libro.save(SALIDA)
    print(f"{len(df)} VUPACI | fórmulas vivas | pesos editables en la hoja Pesos")
    print("Excel:", SALIDA)


def _hoja_pesos(h):
    h["A1"], h["B1"] = "tramo", "peso"
    h["A1"].font = h["B1"].font = NEGRITA
    filas = [("a — llegados 2006-2011", 0.00),
             ("b — llegados 2001-2005", 0.33),
             ("c — llegados 1991-2000", 0.66),
             ("d — llegados antes de 1991", 1.00)]
    for i, (nom, p) in enumerate(filas, start=2):
        h.cell(row=i, column=1, value=nom)
        c = h.cell(row=i, column=2, value=p)
        c.font = NEGRITA
        c.number_format = "0.00"
    h["A7"] = ("Cambiá cualquiera de estos cuatro pesos y la hoja Cálculo "
               "se recalcula sola.")
    h["A7"].font = GRIS
    h.column_dimensions["A"].width = 34
    h.column_dimensions["B"].width = 10


def _hoja_calculo(h, df):
    cabeceras = (["VUPACI", "n secciones"] + CAT
                 + ["Total INE", "a 2006-2011", "b 2001-2005", "c 1991-2000",
                    "d antes 1991", "suma tramos", "p_a", "p_b", "p_c", "p_d",
                    "ÍNDICE DE ARRAIGO"])
    for j, nom in enumerate(cabeceras, start=1):
        c = h.cell(row=1, column=j, value=nom)
        c.font = NEGRITA
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.fill = CABECERA

    L = get_column_letter
    c_2011, c_2006 = L(3), L(8)              # las seis columnas de años sueltos
    c_b, c_c = L(9), L(10)                   # 2001-2005 y 1991-2000
    c_d1, c_d3 = L(11), L(13)                # 1981-1990 .. antes de 1971
    c_tot = L(14)
    c_a, c_bb, c_cc, c_dd = L(15), L(16), L(17), L(18)
    c_sum = L(19)
    c_pa, c_pb, c_pc, c_pd = L(20), L(21), L(22), L(23)
    c_idx = L(24)

    for i, (v, fila) in enumerate(df.iterrows(), start=2):
        marcadas = str(fila.get("marcadas") or "").split("|")
        ausentes = str(fila.get("ausentes") or "").split("|")
        h.cell(row=i, column=1, value=int(v)).font = NEGRITA
        h.cell(row=i, column=2, value=int(fila["n_secciones"]))
        for j, cat in enumerate(CAT, start=3):
            celda = h.cell(row=i, column=j, value=int(fila[cat]))
            if cat in marcadas:
                celda.font, celda.fill = AMBAR_F, AMBAR_R
            elif cat in ausentes:
                celda.font = GRIS
        h.cell(row=i, column=14, value=int(fila["Total"])).font = GRIS

        h[f"{c_a}{i}"] = f"=SUM({c_2011}{i}:{c_2006}{i})"
        h[f"{c_bb}{i}"] = f"={c_b}{i}"
        h[f"{c_cc}{i}"] = f"={c_c}{i}"
        h[f"{c_dd}{i}"] = f"=SUM({c_d1}{i}:{c_d3}{i})"
        h[f"{c_sum}{i}"] = f"=SUM({c_a}{i}:{c_dd}{i})"
        for pc, tc in ((c_pa, c_a), (c_pb, c_bb), (c_pc, c_cc), (c_pd, c_dd)):
            h[f"{pc}{i}"] = f"={tc}{i}/${c_sum}${i}".replace(f"${c_sum}${i}",
                                                             f"{c_sum}{i}")
            h[f"{pc}{i}"].number_format = "0.000"
        h[f"{c_idx}{i}"] = (f"=({c_pa}{i}*Pesos!$B$2+{c_pb}{i}*Pesos!$B$3"
                            f"+{c_pc}{i}*Pesos!$B$4+{c_pd}{i}*Pesos!$B$5)*100")
        h[f"{c_idx}{i}"].font = NEGRITA
        h[f"{c_idx}{i}"].number_format = "0.00"

    h.column_dimensions["A"].width = 9
    h.column_dimensions["B"].width = 10
    for j in range(3, len(cabeceras) + 1):
        h.column_dimensions[get_column_letter(j)].width = 12
    h.column_dimensions[c_idx].width = 15
    h.freeze_panes = "C2"


def _hoja_metodo(h):
    lineas = [
        ("Índice de arraigo 2011 — qué hace cada columna", True),
        ("", False),
        ("Las once primeras columnas son el dato del INE tal cual: personas "
         "residentes en viviendas principales por año de llegada a la "
         "vivienda, censo de 2011, agregadas por VUPACI.", False),
        ("", False),
        ("a, b, c, d son los cuatro tramos de la recodificación:", True),
        ("   a = 2011+2010+2009+2008+2007+2006", False),
        ("   b = 2001-2005", False),
        ("   c = 1991-2000", False),
        ("   d = 1981-1990 + 1971-1980 + antes de 1971", False),
        ("", False),
        ("p_a … p_d son la proporción de cada tramo sobre la suma de los "
         "cuatro, y el índice es:", False),
        ("   índice = (p_a·peso_a + p_b·peso_b + p_c·peso_c + p_d·peso_d) · 100", True),
        ("", False),
        ("Los pesos viven en la hoja Pesos (0 / 0,33 / 0,66 / 1). Cambiarlos "
         "recalcula las 489 filas.", False),
        ("", False),
        ("Decisiones tomadas sobre el correo:", True),
        ("· Los términos son proporciones, no recuentos: sólo así el ·100 da "
         "un índice de 0 a 100. Con recuentos daría millones.", False),
        ("· Los tramos del correo se pisaban («2011-2005» y también "
         "«2001-2005»; «antes de 2000» con «1991-2000» aparte). Se usó la "
         "lectura que no duplica ni deja huecos.", False),
        ("· El denominador es la suma de los cuatro tramos, no la columna "
         "Total INE: difieren hasta en 15 personas por el redondeo del INE a "
         "múltiplos de 5, y así las proporciones cierran en 1 exacto. La "
         "columna Total INE queda a la vista, en gris, para comprobarlo.", False),
        ("· Los asteriscos desaparecen al agregar los tramos, tal como pedía "
         "el correo. Las celdas que venían marcadas siguen en ámbar.", False),
        ("", False),
        ("Lectura: 0 = todos llegaron en los últimos seis años; "
         "100 = todos llegaron antes de 1991.", True),
        ("", False),
        ("En gris cursiva, los seis ceros que el INE no publicó por no tener "
         "ningún dato en esa categoría.", False),
    ]
    for i, (txt, negrita) in enumerate(lineas, start=1):
        c = h.cell(row=i, column=1, value=txt)
        if negrita:
            c.font = NEGRITA
    h.column_dimensions["A"].width = 100


if __name__ == "__main__":
    main()
