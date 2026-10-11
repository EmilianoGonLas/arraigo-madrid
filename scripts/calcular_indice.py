# -*- coding: utf-8 -*-
"""
Índice de arraigo (o de permanencia) por VUPACI, censo 2011.

Recodificación propuesta por Pedro Uceda:

    a) llegados en los últimos años   x 0
    b) 2001-2005                      x 0.33
    c) 1991-2000                      x 0.66
    d) el resto, más antiguos         x 1
    índice = (a + b + c + d) * 100

Dos decisiones que el correo deja abiertas y que aquí se resuelven así
(quedan a la vista para poder cambiarlas en una línea):

1. **Son proporciones, no recuentos.** El correo dice «(a+b+c+d)*100», y eso
   sólo da un índice acotado entre 0 y 100 si cada término es la *proporción*
   de personas de ese tramo sobre el total. Con recuentos daría millones.

2. **Los tramos del correo se pisan.** Dice «sumatorio 2011-2005» y luego
   «2001-2005» por separado, y la (d) habla de «antes de 2000» cuando la (c)
   ya toma 1991-2000. Con las categorías que publica el INE, la lectura
   coherente —y la única que no duplica ni deja huecos— es:

       a = 2011+2010+2009+2008+2007+2006
       b = 2001-2005
       c = 1991-2000
       d = 1981-1990 + 1971-1980 + Antes de 1971

El denominador es la **suma de las once categorías**, no la columna `Total`
del INE: por el redondeo a múltiplos de 5 las dos cifras difieren hasta en 15
personas, y usar la suma hace que las cuatro proporciones cierren en 1 exacto.
Ambos totales quedan en la salida para poder comprobarlo.

Sobre los asteriscos, el correo es explícito: al unificar categorías se suman
sin más, que es justo lo que pasa aquí al agregar los tramos.

Salidas:
  salidas/indice_arraigo_2011.csv           una fila por VUPACI
  salidas/indice_arraigo_2011_secciones.csv una fila por sección (para QGIS)
  salidas/Indice de arraigo 2011.xlsx
"""

import os

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from descargar_arraigo import leer_seccionado

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
ENTRADA = os.path.join(RAIZ, "salidas", "arraigo_por_vupaci.csv")
CSV_VUPACI = os.path.join(RAIZ, "salidas", "indice_arraigo_2011.csv")
CSV_SECCION = os.path.join(RAIZ, "salidas", "indice_arraigo_2011_secciones.csv")
XLSX = os.path.join(RAIZ, "salidas", "Indice de arraigo 2011.xlsx")

TRAMOS = {
    "a_2006_2011": (["2011", "2010", "2009", "2008", "2007", "2006"], 0.00),
    "b_2001_2005": (["2001-2005"], 0.33),
    "c_1991_2000": (["1991-2000"], 0.66),
    "d_antes_1991": (["1981-1990", "1971-1980", "Antes de 1971"], 1.00),
}


def calcular(df):
    out = pd.DataFrame(index=df.index)
    for nombre, (cols, _) in TRAMOS.items():
        out[nombre] = df[cols].sum(axis=1)

    out["suma_categorias"] = out[list(TRAMOS)].sum(axis=1)
    out["total_INE"] = df["Total"]

    for nombre, (_, peso) in TRAMOS.items():
        out[f"p_{nombre}"] = out[nombre] / out["suma_categorias"]

    out["indice_arraigo"] = sum(
        out[f"p_{nombre}"] * peso for nombre, (_, peso) in TRAMOS.items()) * 100
    out["indice_arraigo"] = out["indice_arraigo"].round(2)
    return out


def main():
    df = pd.read_csv(ENTRADA).set_index("VUPACI").sort_index()
    ind = calcular(df)

    # comprobación: las cuatro proporciones tienen que cerrar en 1
    cierre = ind[[f"p_{n}" for n in TRAMOS]].sum(axis=1)
    assert ((cierre - 1).abs() < 1e-9).all(), "las proporciones no cierran en 1"

    ind.insert(0, "n_secciones", df["n_secciones"])
    ind.reset_index().to_csv(CSV_VUPACI, index=False)

    # una fila por sección, para cruzar con la cartografía en QGIS
    grupos = leer_seccionado()
    filas = [{"SECCION": s, "VUPACI": v,
              "indice_arraigo": ind.loc[v, "indice_arraigo"]}
             for v in sorted(grupos) for s in grupos[v]]
    pd.DataFrame(filas).to_csv(CSV_SECCION, index=False)

    with pd.ExcelWriter(XLSX, engine="openpyxl") as xls:
        ind.reset_index().to_excel(xls, sheet_name="Índice", index=False)
        _hoja_metodo(xls.book)
        _formatear(xls.book)

    q = ind["indice_arraigo"].describe(percentiles=[.25, .5, .75])
    print(f"{len(ind)} VUPACI | índice de {q['min']:.1f} a {q['max']:.1f} "
          f"| mediana {q['50%']:.1f} | media {q['mean']:.1f}")
    print("\nmás arraigo:")
    print(ind["indice_arraigo"].nlargest(5).to_string())
    print("\nmenos arraigo:")
    print(ind["indice_arraigo"].nsmallest(5).to_string())
    print(f"\n{CSV_VUPACI}\n{CSV_SECCION}\n{XLSX}")


def _hoja_metodo(libro):
    hoja = libro.create_sheet("Método")
    lineas = [
        ("Índice de arraigo 2011 — cómo está calculado", True),
        ("", False),
        ("tramo", True),
        ("a = 2011+2010+2009+2008+2007+2006     peso 0", False),
        ("b = 2001-2005                          peso 0,33", False),
        ("c = 1991-2000                          peso 0,66", False),
        ("d = 1981-1990 + 1971-1980 + antes de 1971   peso 1", False),
        ("", False),
        ("índice = (pa*0 + pb*0,33 + pc*0,66 + pd*1) * 100", True),
        ("donde pX es la proporción del tramo sobre la suma de las once "
         "categorías.", False),
        ("", False),
        ("Decisiones tomadas sobre el correo:", True),
        ("· Los términos son proporciones, no recuentos: sólo así el *100 da "
         "un índice de 0 a 100.", False),
        ("· Los tramos del correo se pisaban (decía «2011-2005» y también "
         "«2001-2005», y «antes de 2000» con «1991-2000» aparte). Se usó la "
         "lectura que no duplica ni deja huecos.", False),
        ("· El denominador es la suma de las once categorías, no la columna "
         "Total del INE: difieren hasta en 15 personas por el redondeo, y así "
         "las proporciones cierran en 1 exacto. Las dos cifras están en la "
         "hoja Índice.", False),
        ("· Los asteriscos desaparecen al agregar, tal como pedía el correo.", False),
        ("", False),
        ("Lectura: 0 = todos llegaron en los últimos seis años; "
         "100 = todos llegaron antes de 1991.", True),
    ]
    for i, (txt, negrita) in enumerate(lineas, start=1):
        c = hoja.cell(row=i, column=1, value=txt)
        if negrita:
            c.font = Font(bold=True)
    hoja.column_dimensions["A"].width = 105


def _formatear(libro):
    hoja = libro["Índice"]
    for c in hoja[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for j in range(1, hoja.max_column + 1):
        hoja.column_dimensions[get_column_letter(j)].width = 14
    col = [c.value for c in hoja[1]].index("indice_arraigo") + 1
    for i in range(2, hoja.max_row + 1):
        hoja.cell(row=i, column=col).font = Font(bold=True)
    for j, nombre in enumerate([c.value for c in hoja[1]], start=1):
        if str(nombre).startswith("p_"):
            for i in range(2, hoja.max_row + 1):
                hoja.cell(row=i, column=j).number_format = "0.000"
    hoja.freeze_panes = "B2"


if __name__ == "__main__":
    main()
