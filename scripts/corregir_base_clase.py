# -*- coding: utf-8 -*-
"""
Reescribe la base unificada de la clase ("Datos 2011.xlsx") con los valores
verificados contra el INE, manteniendo su misma disposición.

Correcciones aplicadas (ver hoja "Cambios"):
  - VUPACI 84, 321, 337 y 473: cifras que no coincidían con el INE.
  - la fila 84 del archivo original traía además cuatro celdas guardadas como
    decimales (3.41 por 3.410, etc.): Excel se comió el separador de miles.
  - todas las celdas quedan como número, no como texto con asterisco, para que
    las sumas y el índice de arraigo se puedan calcular directamente.
"""

import os

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
NUESTRO = os.path.join(RAIZ, "salidas", "arraigo_por_vupaci.csv")
CLASE = os.path.join(RAIZ, "datos", "originales", "Datos 2011.xlsx")
SALIDA = os.path.join(RAIZ, "salidas", "Datos 2011 - corregido.xlsx")

CAT = ["2011", "2010", "2009", "2008", "2007", "2006",
       "2001-2005", "1991-2000", "1981-1990", "1971-1980", "Antes de 1971"]

CAMBIOS = {
    84: "La base de la clase daba 15.985; el INE da 14.455. Además cuatro "
        "celdas estaban guardadas como decimales (3.41 en vez de 3.410).",
    321: "La base de la clase daba 5.295; el INE da 14.320. Faltaban secciones "
         "en la selección.",
    337: "La base de la clase daba 10.275; el INE da 13.470. Faltaban secciones "
         "en la selección.",
    473: "La base de la clase daba 14.030; el INE da 14.390. Faltaban secciones "
         "en la selección.",
}


def main():
    nos = pd.read_csv(NUESTRO).set_index("VUPACI").sort_index()

    libro = pd.ExcelWriter(SALIDA, engine="openpyxl")
    tabla = nos[["Total"] + CAT].copy()
    tabla.insert(0, "VUPACI", tabla.index)
    tabla.columns = ["VUPACI", "Año de llegada a la vivienda y totales"] + CAT
    tabla.to_excel(libro, sheet_name="Finales", index=False)

    hoja = libro.book["Finales"]
    negrita = Font(bold=True)
    fuente_ast = Font(color="B45309", bold=True)
    relleno_ast = PatternFill("solid", fgColor="FEF3C7")
    fuente_cero = Font(color="6B7280", italic=True)

    for c in hoja[1]:
        c.font = negrita
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    col = {c: j + 3 for j, c in enumerate(CAT)}       # A=VUPACI, B=total
    n_ast = n_cero = 0
    for i, v in enumerate(tabla["VUPACI"], start=2):
        for cat in str(nos.loc[v, "marcadas"] or "").split("|"):
            if cat in col:
                celda = hoja.cell(row=i, column=col[cat])
                celda.font, celda.fill = fuente_ast, relleno_ast
                celda.comment = Comment(
                    "El INE marca esta cifra con asterisco: por debajo del "
                    "umbral de fiabilidad de la muestra. Al unificar "
                    "categorías para el índice, se suma sin más.", "Control")
                n_ast += 1
        for cat in str(nos.loc[v, "ausentes"] or "").split("|"):
            if cat in col:
                celda = hoja.cell(row=i, column=col[cat])
                celda.font = fuente_cero
                celda.comment = Comment(
                    "El INE no publicó esta categoría para este VUPACI: es un "
                    "cero (el total cuadra con la suma de las demás).",
                    "Control")
                n_cero += 1

    hoja.column_dimensions["A"].width = 9
    hoja.column_dimensions["B"].width = 16
    for cat, j in col.items():
        hoja.column_dimensions[get_column_letter(j)].width = 11
    hoja.freeze_panes = "C2"

    cambios = pd.DataFrame(
        [{"VUPACI": v,
          "total en la base de la clase": t,
          "total corregido (INE)": int(nos.loc[v, "Total"]),
          "qué pasaba": CAMBIOS[v]}
         for v, t in [(84, 15985), (321, 5295), (337, 10275), (473, 14030)]])
    cambios.to_excel(libro, sheet_name="Cambios", index=False)
    hc = libro.book["Cambios"]
    for c in hc[1]:
        c.font = negrita
    hc.column_dimensions["A"].width = 9
    hc.column_dimensions["B"].width = 26
    hc.column_dimensions["C"].width = 22
    hc.column_dimensions["D"].width = 90

    libro.close()
    print(f"489 VUPACI | 4 filas corregidas | {n_ast} celdas con asterisco "
          f"marcadas | {n_cero} ceros no publicados")
    print("Excel:", SALIDA)


if __name__ == "__main__":
    main()
